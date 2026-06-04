# -*- coding: utf-8 -*-
"""Build the academic technical-proposal DOCX.
Cover (SUST 光电学院 style) + one-page TOC + body with black IEEE/CCF 3-level
headings + three-line tables + references on a separate page.
Fonts: CN=等线 (cover uses 仿宋/黑体), EN=Times New Roman, code=Consolas.
"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

BLACK = RGBColor(0, 0, 0)
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'assets')


def set_run_font(run, cn='等线', en='Times New Roman', size=12, bold=False,
                 color=BLACK):
    run.font.name = en
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), en)
    rFonts.set(qn('w:hAnsi'), en)
    rFonts.set(qn('w:eastAsia'), cn)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _mk_border(name, sz, val='single', color='000000'):
    el = OxmlElement('w:' + name)
    el.set(qn('w:val'), val)
    el.set(qn('w:sz'), str(sz))      # eighths of a point: 12=1.5pt, 6=0.75pt
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), color)
    return el


def set_three_line(table):
    """Three-line (三线) table: thick top+bottom on table, thin under header."""
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    b.append(_mk_border('top', 12))
    b.append(_mk_border('bottom', 12))
    for n in ('left', 'right', 'insideH', 'insideV'):
        b.append(_mk_border(n, 0, 'none'))
    tblPr.append(b)
    for cell in table.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcB = OxmlElement('w:tcBorders')
        tcB.append(_mk_border('bottom', 6))
        tcPr.append(tcB)


def fill_cell(cell, text, header=False, align='left', size=10.5):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = {'left': WD_ALIGN_PARAGRAPH.LEFT,
                   'center': WD_ALIGN_PARAGRAPH.CENTER,
                   'right': WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=header)


def add_table(doc, headers, rows, aligns=None, caption=None, header_center=True,
              trail=True):
    """three-line table; row/col headers centered. aligns: per-col body align."""
    if caption:
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_before = Pt(6)
        cp.paragraph_format.space_after = Pt(2)
        set_run_font(cp.add_run(caption), size=10.5, bold=True)
    ncol = len(headers)
    aligns = aligns or ['left'] * ncol
    t = doc.add_table(rows=1, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(headers):
        fill_cell(t.rows[0].cells[j], h, header=True, align='center')
    for r in rows:
        cells = t.add_row().cells
        for j, v in enumerate(r):
            # first column acts as a row header -> center it too
            a = 'center' if (header_center and j == 0) else aligns[j]
            fill_cell(cells[j], str(v), align=a)
    set_three_line(t)
    if trail:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def add_heading(doc, text, level):
    """Black IEEE/CCF-style 3-level heading."""
    sizes = {1: 14, 2: 12.5, 3: 12}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    set_run_font(run, cn='黑体', size=sizes[level], bold=True, color=BLACK)
    # outline level so it shows in the TOC field
    pPr = p._p.get_or_add_pPr()
    ol = OxmlElement('w:outlineLvl')
    ol.set(qn('w:val'), str(level - 1))
    pPr.append(ol)
    return p


def add_body(doc, text, size=12, first_indent=True, align='justify'):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.alignment = (WD_ALIGN_PARAGRAPH.JUSTIFY if align == 'justify'
                   else WD_ALIGN_PARAGRAPH.LEFT)
    set_run_font(p.add_run(text), size=size)
    return p


def add_code_block(doc, code):
    for line in code.strip('\n').split('\n'):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(18)
        set_run_font(p.add_run(line if line else ' '),
                     cn='Consolas', en='Consolas', size=9.5)


def add_figure(doc, filename, caption, width_in=5.6):
    """Insert a centered figure with a centered caption below it."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(os.path.join(ASSETS, filename), width=Inches(width_in))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    set_run_font(cap.add_run(caption), size=10.5, bold=True)


def add_keywords(doc, label, words, label_cn='黑体'):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(4)
    set_run_font(p.add_run(label), cn=label_cn, size=12, bold=True)
    set_run_font(p.add_run(words), size=12)



def add_toc_field(doc):
    """Insert a Word TOC field (levels 1-3); user presses F9 / it auto-updates."""
    p = doc.add_paragraph()
    run = p.add_run()
    fldBegin = OxmlElement('w:fldChar'); fldBegin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldSep = OxmlElement('w:fldChar'); fldSep.set(qn('w:fldCharType'), 'separate')
    hint = OxmlElement('w:t'); hint.text = '右键“更新域”生成目录…'
    fldEnd = OxmlElement('w:fldChar'); fldEnd.set(qn('w:fldCharType'), 'end')
    for e in (fldBegin, instr, fldSep, hint, fldEnd):
        run._r.append(e)


def blank(doc, n=1, size=12):
    for _ in range(n):
        p = doc.add_paragraph()
        set_run_font(p.add_run(''), size=size)


def page_break(doc):
    from docx.enum.text import WD_BREAK
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_cover(doc):
    blank(doc, 1)
    # school line — 黑体 16pt centered
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('上海理工大学光电信息与计算机工程学院'),
                 cn='黑体', size=16, bold=True)
    blank(doc, 2)
    # main title — 仿宋 26pt centered
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('《网络监督的细粒度图像识别》'),
                 cn='仿宋', size=26, bold=True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('技术方案'), cn='仿宋', size=26, bold=True)
    blank(doc, 1)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('具身智能机器人挑战赛'), cn='仿宋', size=15, bold=False)
    blank(doc, 2)

    # info fields — labels 仿宋 15.5pt
    def field(label, value):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(f'{label}    {value}')
        set_run_font(run, cn='仿宋', size=15.5)

    field('专　　业', '计算机科学与技术')
    field('年　　级', '2024 级')
    field('指导教师', '王永雄 教授')
    blank(doc, 1)
    members = [
        ('龚琳茜', '2435052601', '算法工程师'),
        ('姚林青', '2435062306', '数据分析师'),
        ('乔蓬旭', '2435053604', '训练工程师'),
        ('陈嘉豪', '2435055208', '项目管理 / 工程'),
    ]
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for j, h in enumerate(('姓　名', '学　号', '分　工')):
        fill_cell(t.rows[0].cells[j], h, header=True, align='center', size=12)
    for nm, sid, role in members:
        cells = t.add_row().cells
        fill_cell(cells[0], nm, align='center', size=12)
        fill_cell(cells[1], sid, align='center', size=12)
        fill_cell(cells[2], role, align='center', size=12)
    blank(doc, 2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('2026 年 6 月'), cn='仿宋', size=15.5)
    page_break(doc)


def setup_doc():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.8)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '等线')
    return doc


def add_page_numbers(doc):
    """Footer page number, centered."""
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = 'PAGE'
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    for x in (b, i, e):
        run._r.append(x)
    set_run_font(run, size=10.5)


# ----------------------------------------------------------------------------
PLACEHOLDER = '__BODY_GOES_HERE__'


def build_body(doc):
    # ===== 摘要 =====
    add_heading(doc, '摘要', 1)
    add_body(doc, '本方案针对“网络监督的细粒度图像识别”赛题，在不使用预训练模型、不引入'
             '额外数据、不进行多模型集成的约束下，研究 ImageNet-1K 上 ResNet-50 从零训练的'
             '噪声鲁棒方案。我们首先通过数据扫描与损失消融判断本数据集属于弱噪声场景'
             '（整体噪声率约 5–15%），进而确立“保留数据覆盖度 + 温和标签空间增强”的主'
             '路线，而非激进样本清洗或高噪声鲁棒损失。最终配方为完整训练集 + 交叉熵/标签'
             '平滑 + Mixup(α=0.1)/CutMix(α=1.0) + raw/EMA 双评估选优，在 100 epoch、单卡 '
             'RTX 4090D 预算下达到验证集 Top-1 77.57%；叠加单模型水平翻转测试时增强'
             '（HFlip TTA）后达到 77.99%、Top-5 93.97%。我们同时报告四组负结果（数据清洗、'
             '强增广、SCE/GCE 鲁棒损失）与三个训练器缺陷的修复，用以支撑“工程严谨性优先于'
             '损失函数堆砌”的核心判断。')
    add_keywords(doc, '关键词：', '图像分类；网络监督学习；噪声标签；ResNet-50；'
                 'Mixup；CutMix；指数滑动平均；测试时增强')

    # ===== 一、算法概述 =====
    add_heading(doc, '一、算法概述', 1)
    add_body(doc, '本方案面向含噪网络监督场景下的 ImageNet-1K 图像分类任务。在不使用预训练'
             '模型、不引入额外数据、不进行多模型集成的约束下，我们采用 ResNet-50 从零训练，'
             '以交叉熵、标签平滑、指数滑动平均（EMA）、余弦学习率和轻量级 Mixup/CutMix 为'
             '核心。最终提交路线为：完整训练集 + ResNet-50 + 交叉熵/标签平滑 + Mixup'
             '（α=0.1）+ CutMix（α=1.0）+ raw/EMA 双评估选优 + 单模型水平翻转测试时增强'
             '（HFlip TTA）。')
    add_body(doc, '验证集结果显示，在同一 raw 权重下基础 CE+LS 基线达到 Top-1 77.58%，最终 '
             'full_mixcut 配方 77.57% 与之基本持平；在同一模型 checkpoint 上叠加单模型 '
             'HFlip TTA 后，full_mixcut 达到 Top-1 77.99%、Top-5 93.97%（baseline 同配置为 '
             '77.97%）。方案的核心判断是：本数据集整体噪声较弱，显式'
             '剔除样本和强鲁棒损失不一定优于保留数据覆盖度；温和的标签空间（label-space）'
             '增强更适合作为噪声鲁棒正则。')

    # ===== 二、实现方案 =====
    add_heading(doc, '二、实现方案', 1)
    add_heading(doc, '2.1　赛题约束与设计原则', 2)
    add_body(doc, '赛题提供 ImageNet-1K 训练集、验证集和测试集，规模如表 1 所示。')
    add_table(doc, ['数据集', '数量', '用途'],
              [['Train', '1,281,167', '从零训练'],
               ['Validation', '50,000', '模型选择与消融'],
               ['Test', '100,000', '生成最终预测 CSV']],
              aligns=['center', 'right', 'left'],
              caption='表 1　ImageNet-1K 数据集规模')
    add_body(doc, '关键约束包括：不能使用预训练模型，不能引入外部数据，最终结果只能对应'
             '一个模型。基于上述约束，我们确立四条设计原则：（1）优先建立可复现强基线，'
             'ResNet-50 是从零训练的稳定主干，训练成本和收敛行为可控；（2）噪声鲁棒以温和'
             '正则为主，在弱噪声场景中避免过早依赖复杂鲁棒损失或大规模删样本；（3）所有'
             '收益必须能被消融解释，每条路线保留训练日志、配置和 checkpoint 元数据；'
             '（4）最终提交保持单模型，HFlip TTA 只对同一个 checkpoint 前向两次，不构成'
             '多模型集成。')

    add_heading(doc, '2.2　数据处理与噪声分析', 2)
    add_body(doc, '我们首先对训练集做完整质量扫描。结果显示数据本身高度可用：损坏文件 0 张、'
             '近空白图 0 张、极小图占比不超过 0.05%，整体可用率约 99.95%，几乎没有需要'
             '丢弃的物理坏样本。这意味着主要挑战不在图像质量，而在类别语义歧义和网络标签'
             '噪声。例如 radio、modem、spatula 等类别存在较高视觉混淆或语义混杂，估计整体'
             '噪声率约 5–15%，个别高噪类别可达 60%。类别数量分布存在差异但不构成严重长尾，'
             '因此没有采用重采样或类别平衡损失（class-balanced loss）作为主路线。')
    add_table(doc, ['扫描维度', '结果', '处理'],
              [['损坏文件', '0 张', '无需剔除'],
               ['近空白图', '0 张', '无需剔除'],
               ['极小图（短边过小）', '≤ 0.05%', '保留，RandomResizedCrop 容忍'],
               ['整体可用率', '99.95%', '不做物理删样本'],
               ['估计标签噪声率', '5–15%（弱噪）', '用软监督而非删样本应对']],
              aligns=['left', 'center', 'left'],
              caption='表 2　训练集数据质量扫描')
    add_body(doc, '基于“数据可用、噪声偏弱”的判断，我们对比了两条应对路线：直接删除高损'
             '样本（清洗），或保留全部样本并用软监督降低噪声影响。我们实现并验证了 C1/C2 '
             '清洗路线，包括剔除明确双类冲突样本、同形异义类别样本以及教师模型高损'
             '（teacher-loss）样本。清洗脚本加入了 exclude 文件 sanity check、teacher '
             'checkpoint key 兼容、每类删除上限和 merge 输入校验，保证清洗列表可复现。'
             '但完整训练结果表明，清洗数据的 Top-1 为 76.84%，低于 full baseline（77.58%）'
             '和 full_mixcut（77.57%）。该结果说明：在本任务中，直接删除高损样本可能同时'
             '删除了困难但有效的训练样本，损害数据覆盖度。因此最终选择“保留样本、软化'
             '监督”的主路线。')
    add_heading(doc, '2.3　模型结构', 2)
    add_body(doc, '主模型为标准 ResNet-50，类别数为 1000。选择 ResNet-50 的对比与定位如表 3'
             '所示。最终路线未使用预训练权重，模型输出 1000 类 logits，训练和验证均严格'
             '按照 ImageNet-1K 类别索引进行。')
    add_table(doc, ['模型', '参数量', '优点', '本项目定位'],
              [['StarNet-S2', '约 12M', '训练快', '早期快速筛选'],
               ['ResNet-50', '约 25M', '稳定、可复现、从零训练成熟', '最终提交主模型'],
               ['ConvNeXt-V2-T', '约 28M', '潜在上限更高', '算力充足时的备选']],
              aligns=['center', 'right', 'left', 'left'],
              caption='表 3　候选主干网络对比')

    add_heading(doc, '2.4　训练策略', 2)
    add_body(doc, '基础训练策略如表 4 所示。最终 full_mixcut 配方在此基础上只加入 batch 级'
             '标签空间增强，具体增强配置见表 5。')
    add_table(doc, ['模块', '设置'],
              [['Loss', 'Cross Entropy + Label Smoothing 0.1'],
               ['Optimizer', 'SGD，momentum 0.9，weight decay 1e-4'],
               ['LR schedule', '5 epoch 线性预热 + 余弦衰减'],
               ['Epochs', '100'],
               ['Batch', 'per-step 128，梯度累积 2，等效 batch 256'],
               ['Precision', 'AMP 混合精度'],
               ['Regularization', 'EMA decay 0.9999，grad clip 5.0'],
               ['Hardware', '单卡 RTX 4090D']],
              aligns=['center', 'left'],
              caption='表 4　基础训练策略')
    add_table(doc, ['增强', '参数', '说明'],
              [['RandomResizedCrop', '224', '基础空间增强'],
               ['RandomHorizontalFlip', 'p=0.5', '基础空间增强'],
               ['Mixup', 'α=0.1', '样本与标签线性混合，缓解硬标签噪声'],
               ['CutMix', 'α=1.0', '图像区域级混合，增强局部判别鲁棒性'],
               ['mixup_prob', '0.5', '同时启用时各 50% batch 使用 Mixup / CutMix'],
               ['RandAugment / RandomErasing', '不启用', '100 epoch 预算下强增广收敛不足']],
              aligns=['center', 'center', 'left'],
              caption='表 5　full_mixcut 标签空间增强配置')
    add_body(doc, 'Mixup/CutMix 的 target 采用 one-hot 后混合，再叠加 label smoothing，'
             '将原始 one-hot 硬标签转成软标签，降低噪声标签对单一样本梯度的支配强度。'
             'CutMix 实现中对输入图像使用 clone 后贴 patch，避免 in-place 修改影响后续'
             '数据路径。图 1 给出最终 full_mixcut 配方在 100 epoch 内的训练损失与验证 '
             'Top-1 曲线：训练损失从 6.86 平滑下降至 2.65，验证 Top-1 随余弦学习率稳定'
             '爬升，末段无过拟合回落，最终收敛于 77.57%。')
    add_figure(doc, 'fig1_training_curve.png',
               '图 1　full_mixcut 训练损失与验证 Top-1 曲线（100 epoch，单卡 4090D）')

    add_heading(doc, '2.5　EMA 与双评估模型选择', 2)
    add_body(doc, '早期 baseline 只使用 EMA 模型验证。复盘发现，EMA decay=0.9999 在 100 '
             'epoch 余弦末段可能滞后于 raw model，尤其当学习率快速降低时，raw 权重反而'
             '可能是当前最优。因此 trainer 中实现了 raw + EMA 双评估：每个 epoch 同时'
             '验证 raw model 和 EMA model，分别保存 best_raw.pth、best_ema.pth；全局 '
             'best.pth 保存当轮更高者并记录 best_source；evaluate.py 和 predict.py 根据 '
             'best_source 自动加载 raw 或 EMA 权重。最终 full_mixcut 的最佳 checkpoint '
             '元数据如下：')
    add_code_block(doc, '''checkpoints_full_mixcut/best.pth
epoch        = 100
best_source  = raw
best_acc     = 77.57
best_raw_acc = 77.57
best_ema_acc = 77.438''')
    add_body(doc, '这说明 dual-eval 避免了“raw 已超过 EMA，但推理仍加载 EMA”的精度损失。'
             '图 2 给出训练后段（第 50–100 epoch）raw 与 EMA 验证 Top-1 的对比：前段 EMA '
             '凭借权重平均更平稳、且一度领先，但随着余弦学习率快速衰减，raw 模型在末段'
             '反超 EMA，最终 raw 77.57% 高于 EMA 77.438%。若沿用早期“只验 EMA”的做法，'
             '将损失约 0.13 个百分点。')
    add_figure(doc, 'fig4_raw_ema.png',
               '图 2　训练后段 raw 与 EMA 验证 Top-1 对比（第 50–100 epoch）')

    add_heading(doc, '2.6　推理与提交', 2)
    add_body(doc, '最终推理使用单个 checkpoint，命令如下：')
    add_code_block(doc, '''python scripts/predict.py \\
  --config configs/imagenet_resnet50_full_mixcut.yaml \\
  --checkpoint checkpoints_full_mixcut/best.pth \\
  --tta hflip \\
  --output submit/result.csv''')
    add_body(doc, 'HFlip TTA 流程为：原图前向一次、水平翻转图前向一次，对同一模型的两次 '
             'softmax 结果求和后取 argmax。该方法不使用额外模型，不属于多模型集成。在显存'
             '管理上，10 万张测试图的 logits 张量约为 10 万 × 1000 × 4 字节 ≈ 400MB，'
             '我们将其累加在 CPU 上，避免 GPU 显存溢出。验证集 TTA 结果如表 6 所示，'
             '并与 CE+LS baseline 在同一 raw 权重、同一评估精度下对照。')
    add_table(doc, ['模型', '推理方式', 'Val Top-1', 'Val Top-5'],
              [['baseline (CE+LS)', 'Plain', '77.58', '93.64'],
               ['baseline (CE+LS)', 'HFlip TTA', '77.97', '93.91'],
               ['full_mixcut', 'Plain', '77.57', '93.77'],
               ['full_mixcut', 'HFlip TTA', '77.99', '93.97']],
              aligns=['left', 'center', 'right', 'right'],
              caption='表 6　baseline 与 full_mixcut 的 TTA 对照（raw 权重，fp16）')
    add_body(doc, '表 6 显示：无论加不加 TTA，full_mixcut 与 baseline 均基本持平（Plain '
             '−0.01pp，TTA +0.02pp），差距远小于单次训练的随机波动。HFlip TTA 对两者的'
             '增益相近（约 +0.4pp），说明该增益来自推理端的视角集成，而非某一训练配方'
             '的专属收益。最终提交采用 full_mixcut raw + HFlip TTA（77.99%）。')
    add_body(doc, '最终预测文件为 submit/result.csv，共 100,000 行，已通过格式校验：'
             '每行包含 filename 与四位类别编号。')

    # ===== 三、算法创新点 =====
    add_heading(doc, '三、算法创新点', 1)
    add_heading(doc, '3.1　弱噪声场景下的路线选择', 2)
    add_body(doc, '我们没有简单套用高噪声标签学习中的强鲁棒损失，而是通过实验判断噪声强度'
             '和训练预算。在 40 epoch、弱增广的统一对照下，对称交叉熵（SCE）与广义交叉熵'
             '（GCE）均未超过 CE+LS 基线：SCE(α=0.5,β=0.5) 与 CE 完全持平（66.63%），'
             'SCE(α=0.1,β=1.0) 因 β 过大退化至 60.66%，GCE(q=0.7) 早期梯度被截断、warmup '
             '不足，直接崩盘至 26.01%。具体如表 7 所示。这组结果说明本赛题更接近“弱噪声 + '
             '大规模数据覆盖”场景——SCE/GCE 是为 40% 以上高噪声率设计的，其假设在此不成立。')
    add_table(doc, ['损失函数', '配置', 'Val Top-1', '相对基线'],
              [['CE + LS', 'smoothing=0.1（基线）', '66.63', '—'],
               ['SCE', 'α=0.5, β=0.5', '66.63', '±0.00'],
               ['SCE', 'α=0.1, β=1.0', '60.66', '−5.97'],
               ['GCE', 'q=0.7（崩溃）', '26.01', '−40.62']],
              aligns=['center', 'left', 'right', 'right'],
              caption='表 7　鲁棒损失对照（40 epoch，弱增广，30% 子集）')
    add_heading(doc, '3.2　Mixup/CutMix 作为隐式噪声鲁棒监督', 2)
    add_body(doc, '相比直接删除高损样本，Mixup/CutMix 保留所有训练样本，同时把监督信号从'
             '硬标签变成软标签。它们通过邻域风险最小化（vicinal risk minimization）降低'
             '单个错误标签对参数更新的影响，适合作为网络监督数据的温和鲁棒策略。')
    add_body(doc, '早期快速验证阶段，全量数据 30 epoch 下 mixup-only 已达 69.78%，显著高于'
             '损失对比阶段 30% 子集 40 epoch 的 CE+LS 66.63%；虽然两者训练条件不同，但'
             '趋势已足够明确。最终依据来自 100 epoch 受控消融，且在同一权重（raw）、同一'
             '评估精度下与基线严格对照：不加 TTA 时 full_mixcut（Mixup α=0.1 + CutMix '
             'α=1.0）77.57% 与纯 CE+LS baseline 77.58% 基本持平（−0.01pp）；叠加单模型 '
             'HFlip TTA 后 full_mixcut 77.99% 同样与 baseline 77.97% 持平（+0.02pp）。两组'
             '对比均明显优于全强增广（mixup_v2 + RandAugment/Erasing）的 76.91%（−0.66pp）。'
             '这说明在本含噪数据集上，Mixup/CutMix 作为温和的标签空间正则相对基线不掉点，'
             '且显著优于激进的像素级增广，支撑了我们把算力押向标签空间增强、而非强增广的决策。')
    add_heading(doc, '3.3　raw/EMA 双评估的 checkpoint 选择', 2)
    add_body(doc, 'EMA 通常提高稳定性，但不是所有 epoch 都优于 raw model。最终 full_mixcut '
             '的最优源是 raw（见图 2），这证明双评估不是冗余功能，而是直接影响最终模型'
             '选择的关键环节。')
    add_heading(doc, '3.4　负结果驱动的工程闭环', 2)
    add_body(doc, '项目系统性保留了 mixup_v2 全强增广、C1/C2 清洗、SCE/GCE 损失等负结果。'
             '图 3 汇总了核心消融的验证 Top-1：full_mixcut（77.57%）与叠加 HFlip TTA 后的'
             ' 77.99% 明显高于 cleaned_v1（76.84%）和 full_strong_aug（76.91%）。图 4 进一步'
             '对比三条 100 epoch 路线的收敛曲线，可见 full_strong_aug 因强增广在 100 epoch '
             '预算下前期收敛迟缓、后期才追平，印证了“强增广更适合长 schedule”的判断。这些'
             '负结果帮助我们排除不合适的路线，最终收敛到更简单、可复现、验证更强的方案。')
    add_figure(doc, 'fig2_ablation_bars.png',
               '图 3　核心消融实验验证 Top-1 对比')
    add_figure(doc, 'fig3_convergence.png',
               '图 4　三条 100 epoch 路线验证 Top-1 收敛曲线对比')
    add_heading(doc, '3.5　训练器缺陷的发现与修复', 2)
    add_body(doc, '严谨的归因依赖可信的训练器。在 5 月底的全代码审查中，我们定位并修复了'
             '三个会污染实验结论的缺陷（见表 8）。这三处修复把“看不见的实验变量”变成了'
             '“看得见的方法论”，保证了各路线之间的可比性，以及最终推理加载的是正确权重。')
    add_table(doc, ['缺陷', '影响', '修复'],
              [['标签平滑静默失效', 'Mixup 分支跳过 label smoothing，使 mixup 系列与基线变量不可控，归因不干净', '在 mixup 路径显式注入 mixed×(1−ε)+ε/N'],
               ['cudnn 配置冲突', 'deterministic=True 时 benchmark 实际不生效，100 epoch 慢 5–15%（约 2–7 小时）', '权衡 bit-exact 换速度：deterministic=False、benchmark=True'],
               ['resume 调度器错误', 'cosine 旧 state 替换进新 schedule，导致学习率反向爬升炸毁权重，验证 72.51% → 12.64%', '新增 --resume-mode finetune，仅加载 model+ema，重置调度器']],
              aligns=['left', 'left', 'left'], header_center=False,
              caption='表 8　训练器关键缺陷与修复')

    # ===== 四、问题与思考 =====
    add_heading(doc, '四、问题与思考', 1)
    add_body(doc, '研发过程中遇到的主要问题、观察与处理思路汇总于表 9。')
    add_table(doc, ['问题', '观察', '处理与思考'],
              [['强增广收敛不足', 'Mixup+CutMix+RandAugment+RandomErasing 的 100ep 结果为 76.91，低于 baseline', '该配方更接近长 schedule recipe，100ep 下去掉 RandAugment 和 RandomErasing'],
               ['数据清洗未带来收益', 'cleaned_v1 Top-1 为 76.84', '高损样本可能包含有效困难样本，删除会降低覆盖度'],
               ['鲁棒损失无收益', 'SCE 持平或退化，GCE 崩盘', '弱噪声下 CE+LS 已足够强，复杂损失会增加调参成本'],
               ['EMA 末段滞后', 'full_mixcut raw 77.57，高于 EMA 77.438', '使用 raw/EMA dual-eval 和 best_source 元数据'],
               ['复现交付', '需同时交付训练、验证、预测、Docker 与报告', '以 best.pth 为最终模型，latest.pth 仅用于 resume']],
              aligns=['left', 'left', 'left'], header_center=False,
              caption='表 9　关键问题与处理')
    add_body(doc, '进一步改进方向包括：在不使用额外数据的前提下尝试更长 schedule、轻量 '
             'RandAugment、ConvNeXt 类主干，或对高噪声类别做“软权重”而非直接删样本。')

    # ===== 五、过程进度 =====
    add_heading(doc, '五、过程进度', 1)
    add_body(doc, '项目自数据准备到交付整理的主要阶段与时间节点如表 10 所示。')
    add_table(doc, ['阶段', '时间', '主要工作'],
              [['数据准备', '4/28 – 5/4', '数据下载、目录检查、ImageNet val 结构修复、环境部署'],
               ['Pipeline 搭建', '5/5 – 5/11', 'ResNet/StarNet 训练框架、loss、metrics、checkpoint、predict'],
               ['基线训练', '5/15 – 5/18', 'ResNet-50 100ep baseline，Top-1 77.58'],
               ['鲁棒 loss 消融', '5/19 – 5/25', 'CE/SCE/GCE 对比，确认复杂 loss 不优于 CE+LS'],
               ['增强路线探索', '5/22 – 5/27', 'Mixup/CutMix 短跑与 mixup_v2 强增广长跑'],
               ['代码审查与修复', '5/28 – 5/30', '修复 label smoothing、CutMix clone、resume、dual-eval'],
               ['清洗路线验证', '5/30 – 6/1', 'C1+C2 cleaned_v1，结果 76.84，未作为最终提交'],
               ['最终训练与推理', '5/30 – 6/2', 'full_mixcut 100ep，plain 77.57，HFlip TTA 77.99'],
               ['交付整理', '6/2 起', 'Docker、技术报告、答辩材料、提交文档']],
              aligns=['center', 'center', 'left'],
              caption='表 10　项目过程进度')

    # ===== 六、团队分工 =====
    add_heading(doc, '六、团队分工', 1)
    add_table(doc, ['成员', '学号', '角色', '具体负责内容'],
              [['姚林青', '2435062306', '数据分析师', '数据探索、噪声分析、类别分布统计、可视化'],
               ['乔蓬旭', '2435053604', '训练工程师', 'AutoDL 环境、训练调度、正式实验执行、日志整理'],
               ['龚琳茜', '2435052601', '算法工程师', 'Loss 函数、噪声鲁棒策略、消融实验设计'],
               ['陈嘉豪', '2435055208', '项目管理 / 工程', 'Pipeline 搭建、代码审查、Docker 可复现、报告与答辩']],
              aligns=['center', 'center', 'center', 'left'],
              caption='表 11　团队成员分工', trail=False)

    # ===== 参考文献（单独一页）=====
    refh = add_heading(doc, '参考文献', 1)
    refh.paragraph_format.page_break_before = True
    refs = [
        'He K, Zhang X, Ren S, et al. Deep Residual Learning for Image Recognition[C]. '
        'IEEE Conference on Computer Vision and Pattern Recognition (CVPR), 2016: 770-778.',
        'Zhang H, Cisse M, Dauphin Y N, et al. mixup: Beyond Empirical Risk Minimization[C]. '
        'International Conference on Learning Representations (ICLR), 2018.',
        'Yun S, Han D, Oh S J, et al. CutMix: Regularization Strategy to Train Strong '
        'Classifiers with Localizable Features[C]. IEEE/CVF International Conference on '
        'Computer Vision (ICCV), 2019: 6023-6032.',
        'Szegedy C, Vanhoucke V, Ioffe S, et al. Rethinking the Inception Architecture for '
        'Computer Vision[C]. IEEE Conference on Computer Vision and Pattern Recognition '
        '(CVPR), 2016: 2818-2826.',
        'Wang Y, Ma X, Chen Z, et al. Symmetric Cross Entropy for Robust Learning with Noisy '
        'Labels[C]. IEEE/CVF International Conference on Computer Vision (ICCV), 2019: 322-330.',
        'Zhang Z, Sabuncu M R. Generalized Cross Entropy Loss for Training Deep Neural '
        'Networks with Noisy Labels[C]. Advances in Neural Information Processing Systems '
        '(NeurIPS), 2018: 8778-8788.',
        'Paszke A, Gross S, Massa F, et al. PyTorch: An Imperative Style, High-Performance '
        'Deep Learning Library[C]. Advances in Neural Information Processing Systems '
        '(NeurIPS), 2019: 8024-8035.',
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Pt(20)
        p.paragraph_format.first_line_indent = Pt(-20)
        set_run_font(p.add_run(f'[{i}]　{r}'), size=10.5)


def main():
    doc = setup_doc()
    add_page_numbers(doc)
    build_cover(doc)
    # ----- TOC page -----
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run('目　录'), cn='黑体', size=16, bold=True)
    blank(doc, 1)
    add_toc_field(doc)
    page_break(doc)
    # ----- body -----
    build_body(doc)
    out = r'a:\VScode\Code\Projects\embodied_ai_contest\submit\tech_report\技术方案.docx'
    doc.save(out)
    print('saved', out)


if __name__ == '__main__':
    main()








