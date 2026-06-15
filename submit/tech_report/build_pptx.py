# -*- coding: utf-8 -*-
"""Build the defense PPTX (third backup deliverable).
Design mirrors answer_deck.html: warm paper #fafaf8, ink #0a0a0a, academic
crimson #9E1B1B accent, mono kicker labels, big hero numbers, one idea / slide.
Fonts: CN=微软雅黑, Latin=Segoe UI, mono kicker=Consolas."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# palette
PAPER = RGBColor(0xFA, 0xFA, 0xF8)
INK = RGBColor(0x0A, 0x0A, 0x0A)
ACCENT = RGBColor(0x9E, 0x1B, 0x1B)
BRIGHT = RGBColor(0xD9, 0x45, 0x45)
GREY3 = RGBColor(0x73, 0x73, 0x73)
GREY2 = RGBColor(0xD4, 0xD4, 0xD2)
SEC = RGBColor(0x52, 0x52, 0x52)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

CN = '微软雅黑'
EN = 'Segoe UI'
MONO = 'Consolas'

EMU = 914400
SW, SH = 13.333, 7.5  # 16:9 inches

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def _set_font(run, cn=CN, en=EN):
    run.font.name = en
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', cn)


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = PAPER
    bg.line.fill.background()
    bg.shadow.inherit = False
    s.shapes._spTree.remove(bg._element)
    s.shapes._spTree.insert(2, bg._element)
    return s


def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        space_after=4, line=1.0):
    """runs: list of dicts {t, size, color, bold, cn, en, mono, tracking}."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(runs, dict):
        runs = [runs]
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(space_after)
    p.line_spacing = line
    for r in runs:
        run = p.add_run()
        run.text = r['t']
        run.font.size = Pt(r.get('size', 18))
        run.font.bold = r.get('bold', False)
        run.font.color.rgb = r.get('color', INK)
        mono = r.get('mono', False)
        _set_font(run, cn=r.get('cn', CN), en=(MONO if mono else r.get('en', EN)))
        if r.get('tracking'):
            rPr = run._r.get_or_add_rPr()
            rPr.set('spc', str(int(r['tracking'] * 100)))
    return tb


def rect(s, x, y, w, h, color, line=None, lw=1.0):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y),
                            Inches(w), Inches(h))
    if color is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    return sp


def hline(s, x, y, w, color=GREY2, weight=1.0):
    ln = s.shapes.add_connector(2, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = color
    ln.line.width = Pt(weight)
    return ln


def kicker(s, x, y, text, color=ACCENT, size=11):
    """mono uppercase eyebrow label."""
    return txt(s, x, y, 11, 0.3,
               {'t': text, 'size': size, 'color': color, 'mono': True,
                'tracking': 1.5, 'bold': True})


def page_num(s, n):
    txt(s, SW - 1.5, SH - 0.55, 1.1, 0.3,
        {'t': f'{n:02d} / 10', 'size': 10, 'color': GREY3, 'mono': True,
         'tracking': 1.0}, align=PP_ALIGN.RIGHT)


def stat_block(s, x, y, label, value, color=INK, w=2.2):
    txt(s, x, y, w, 0.3, {'t': label, 'size': 10, 'color': GREY3, 'mono': True,
                          'tracking': 1.0})
    txt(s, x, y + 0.28, w, 0.9, {'t': value, 'size': 32, 'color': color,
                                 'bold': True, 'en': EN})


# ============================ SLIDE 1 · COVER ============================
def s1():
    s = slide()
    rect(s, 0, 0, 0.22, SH, ACCENT)  # left accent bar
    kicker(s, 0.7, 0.6, 'USST · 嵌入式 AI 大赛 · WEBLY-SUPERVISED FGIR')
    txt(s, 0.68, 1.05, 12, 0.35,
        {'t': 'WEBLY-SUPERVISED FINE-GRAINED IMAGE RECOGNITION',
         'size': 12, 'color': GREY3, 'mono': True, 'tracking': 1.2})
    txt(s, 0.66, 1.5, 12, 1.4,
        {'t': '网络监督的细粒度图像识别', 'size': 46, 'color': INK, 'bold': True})
    txt(s, 0.68, 2.75, 12, 0.5, [
        {'t': 'ResNet-50 from scratch ', 'size': 18, 'color': SEC},
        {'t': '+ Mixup/CutMix 单模型方案', 'size': 18, 'color': ACCENT, 'bold': True}])
    txt(s, 0.68, 3.25, 12, 0.35,
        {'t': 'SINGLE 4090D · FROM-SCRATCH · 100 EP', 'size': 11,
         'color': GREY3, 'mono': True, 'tracking': 1.2})
    # hero stats
    hline(s, 0.7, 4.2, 7.0, GREY2, 1.0)
    stat_block(s, 0.7, 4.45, 'BASELINE', '77.58%', INK)
    stat_block(s, 3.1, 4.45, 'FINAL · HFLIP TTA', '77.99%', ACCENT)
    stat_block(s, 5.5, 4.45, 'TRAIN DATA', '128 万', INK)
    # members
    hline(s, 0.7, 5.9, 11.9, GREY2, 1.0)
    members = [('D · 项目 / 工程', '陈嘉豪'), ('A · 训练', '乔蓬旭'),
               ('B · 数据', '姚林青'), ('C · 算法', '龚琳茜')]
    for i, (role, name) in enumerate(members):
        mx = 0.7 + i * 3.0
        txt(s, mx, 6.1, 2.8, 0.3, {'t': role, 'size': 10, 'color': ACCENT,
                                   'mono': True, 'tracking': 0.5})
        txt(s, mx, 6.38, 2.8, 0.4, {'t': name, 'size': 16, 'color': INK,
                                    'bold': True})
    txt(s, 0.7, 6.95, 6, 0.3, {'t': '指导教师 · 王永雄 教授   |   2026.06',
                               'size': 11, 'color': GREY3})
    page_num(s, 1)


def header(s, n, kick, title, sub=None):
    """standard content-slide header: top rule + kicker + title."""
    rect(s, 0, 0, 0.22, SH, ACCENT)
    kicker(s, 0.7, 0.55, kick)
    txt(s, 0.66, 0.95, 12, 1.0, {'t': title, 'size': 30, 'color': INK,
                                 'bold': True})
    if sub:
        txt(s, 0.68, 1.75, 12, 0.4, {'t': sub, 'size': 13, 'color': SEC})
    page_num(s, n)


# ===================== SLIDE 2 · 三大挑战 =====================
def s2():
    s = slide()
    header(s, 2, 'PROBLEM · 三大挑战', '三个挑战，一个判断',
           '弱噪场景的最大公约数 = 数据增广，不是噪声鲁棒损失')
    cards = [
        ('01', 'LABEL · 标签层', '标签噪声 Label Noise',
         '整体噪声率 5–15%；radio / modem / spatula 等高噪类别可达 60%，标签歧义严重',
         '5–15%', 'OVERALL'),
        ('02', 'CLASS · 类别层', '长尾分布 Long-Tail',
         '类别数量分布存在差异但不构成严重长尾；未把重采样或 class-balanced loss 作为主路线',
         '1.78×', 'HEAD / TAIL'),
        ('03', 'DOMAIN · 域偏移', '数据偏差 Domain Shift',
         'web 图像与标准 ImageNet val 域分布存在偏移，弱噪场景下数据增广收益更稳健',
         '128万→5万', 'TRAIN / VAL'),
    ]
    cw, gap, x0, y0, ch = 3.7, 0.3, 0.7, 2.5, 4.0
    for i, (no, tag, name, desc, big, blab) in enumerate(cards):
        cx = x0 + i * (cw + gap)
        rect(s, cx, y0, cw, ch, WHITE, line=GREY2, lw=0.75)
        rect(s, cx, y0, cw, 0.08, ACCENT)
        txt(s, cx + 0.3, y0 + 0.3, 1, 0.5, {'t': no, 'size': 26, 'color': GREY2,
                                            'bold': True, 'mono': True})
        txt(s, cx + 0.3, y0 + 0.95, cw - 0.6, 0.3, {'t': tag, 'size': 9.5,
            'color': ACCENT, 'mono': True, 'tracking': 0.8})
        txt(s, cx + 0.3, y0 + 1.25, cw - 0.6, 0.5, {'t': name, 'size': 16,
            'color': INK, 'bold': True})
        txt(s, cx + 0.3, y0 + 1.85, cw - 0.6, 1.4, {'t': desc, 'size': 11.5,
            'color': SEC}, line=1.25)
        hline(s, cx + 0.3, y0 + 3.25, cw - 0.6, GREY2, 0.75)
        txt(s, cx + 0.3, y0 + 3.35, cw - 0.6, 0.5, {'t': big, 'size': 22,
            'color': ACCENT, 'bold': True})
        txt(s, cx + 0.3, y0 + 3.78, cw - 0.6, 0.3, {'t': blab, 'size': 9,
            'color': GREY3, 'mono': True, 'tracking': 0.8})


# ===================== SLIDE 3 · Pipeline / 硬件 / Timeline =====================
def s3():
    s = slide()
    header(s, 3, 'FRAMEWORK · 四阶段流水线', '从 web 数据到 result.csv',
           'FOUR-STAGE PIPELINE · SINGLE GPU FROM-SCRATCH')
    stages = [
        ('01', '数据准备', 'val 重建 · 1000 类校验 · symlink 修复'),
        ('02', '训练 · ResNet-50', 'Mixup/CutMix · 100 ep · EMA · AMP · accum=2'),
        ('03', '推理 · 单模型 HFlip TTA', '+0.42 pp Top-1 · one checkpoint'),
        ('04', '提交', 'result.csv · Dockerfile · 技术方案'),
    ]
    y0 = 2.45
    for i, (no, name, desc) in enumerate(stages):
        sy = y0 + i * 0.95
        rect(s, 0.7, sy, 0.62, 0.62, ACCENT)
        txt(s, 0.7, sy + 0.07, 0.62, 0.5, {'t': no, 'size': 18, 'color': WHITE,
            'bold': True, 'mono': True}, align=PP_ALIGN.CENTER)
        txt(s, 1.55, sy + 0.02, 5.4, 0.4, {'t': name, 'size': 15, 'color': INK,
            'bold': True})
        txt(s, 1.55, sy + 0.42, 5.6, 0.4, {'t': desc, 'size': 11, 'color': SEC})
        if i < 3:
            ln = s.shapes.add_connector(2, Inches(1.01), Inches(sy + 0.62),
                                        Inches(1.01), Inches(sy + 0.95))
            ln.line.color.rgb = GREY2; ln.line.width = Pt(1.5)
    # right rail: hardware + key dates
    rx = 7.7
    rect(s, rx, 2.45, 4.9, 4.3, WHITE, line=GREY2, lw=0.75)
    txt(s, rx + 0.35, 2.7, 4, 0.3, {'t': 'HARDWARE · 单卡', 'size': 10,
        'color': ACCENT, 'mono': True, 'tracking': 1.0})
    txt(s, rx + 0.35, 3.0, 4.2, 0.6, {'t': 'RTX 4090D', 'size': 26,
        'color': INK, 'bold': True})
    hline(s, rx + 0.35, 3.75, 4.2, GREY2, 0.75)
    pairs = [('TRAIN IMAGES', '128 万'), ('CLASSES', '1 000'),
             ('EPOCHS', '100'), ('EFFECTIVE BATCH', '256')]
    for i, (l, v) in enumerate(pairs):
        px = rx + 0.35 + (i % 2) * 2.25
        py = 3.95 + (i // 2) * 1.0
        txt(s, px, py, 2.1, 0.3, {'t': l, 'size': 9, 'color': GREY3,
            'mono': True, 'tracking': 0.6})
        txt(s, px, py + 0.26, 2.1, 0.5, {'t': v, 'size': 20, 'color': INK,
            'bold': True})
    txt(s, rx + 0.35, 6.1, 4.2, 0.5, [
        {'t': '提交 6/6', 'size': 13, 'color': ACCENT, 'bold': True},
        {'t': '    答辩 6/14', 'size': 13, 'color': SEC}])


# ===================== SLIDE 4 · 数据质量 × Recipe =====================
def s4():
    s = slide()
    header(s, 4, 'DATA · 质量 × 增广 RECIPE', '瓶颈在标签，不在图像')
    # panel A
    ax, ay, aw, ah = 0.7, 2.4, 5.75, 3.6
    rect(s, ax, ay, aw, ah, WHITE, line=GREY2, lw=0.75)
    txt(s, ax + 0.35, ay + 0.3, aw - 0.7, 0.3, {'t': 'A · 数据质量', 'size': 11,
        'color': ACCENT, 'mono': True, 'tracking': 1.0})
    txt(s, ax + 0.35, ay + 0.62, aw - 0.7, 0.6, [
        {'t': '128 万张  ', 'size': 22, 'color': INK, 'bold': True},
        {'t': '0 张报废', 'size': 16, 'color': ACCENT, 'bold': True}])
    for i, (l, v) in enumerate([('损坏文件', '0 张'), ('近空白图', '0 张'),
                                ('极小图', '≤ 0.05%')]):
        ly = ay + 1.35 + i * 0.42
        txt(s, ax + 0.35, ly, 3, 0.3, {'t': f'{i+1:02d}  {l}', 'size': 12,
            'color': SEC})
        txt(s, ax + 3.4, ly, 2, 0.3, {'t': v, 'size': 12, 'color': INK,
            'bold': True})
    hline(s, ax + 0.35, ay + 2.75, aw - 0.7, GREY2, 0.75)
    txt(s, ax + 0.35, ay + 2.9, 2.5, 0.5, [
        {'t': '99.95%', 'size': 22, 'color': ACCENT, 'bold': True}])
    txt(s, ax + 0.35, ay + 3.32, 3, 0.25, {'t': 'USABLE', 'size': 9,
        'color': GREY3, 'mono': True, 'tracking': 1.0})
    # panel B
    bx = 6.85
    rect(s, bx, ay, aw, ah, WHITE, line=GREY2, lw=0.75)
    txt(s, bx + 0.35, ay + 0.3, aw - 0.7, 0.3, {'t': 'B · 增广 RECIPE', 'size': 11,
        'color': ACCENT, 'mono': True, 'tracking': 1.0})
    txt(s, bx + 0.35, ay + 0.62, aw - 0.7, 0.6, [
        {'t': '77.99%  ', 'size': 22, 'color': ACCENT, 'bold': True},
        {'t': '最终提交 · raw + HFlip TTA', 'size': 12, 'color': SEC}])
    for i, (l, v) in enumerate([('基础', 'RRC(224) + HFlip'),
                                ('取舍', '不启用 RandAug / Erasing'),
                                ('正则', 'Mixup α=0.1 · CutMix α=1.0 · p=0.5')]):
        ly = ay + 1.35 + i * 0.42
        txt(s, bx + 0.35, ly, 1.1, 0.3, {'t': l, 'size': 11, 'color': ACCENT,
            'bold': True})
        txt(s, bx + 1.45, ly, 4.0, 0.3, {'t': v, 'size': 11, 'color': SEC})
    hline(s, bx + 0.35, ay + 2.75, aw - 0.7, GREY2, 0.75)
    txt(s, bx + 0.35, ay + 2.9, 5, 0.44, [
        {'t': '不加 TTA  ', 'size': 11, 'color': GREY3, 'mono': True},
        {'t': 'mixcut 77.57 ≈ base 77.58 > 强增广 76.91', 'size': 11, 'color': SEC}])
    txt(s, bx + 0.35, ay + 3.28, 5, 0.25, [
        {'t': '+HFlip TTA  ', 'size': 11, 'color': ACCENT, 'mono': True, 'bold': True},
        {'t': 'mixcut 77.99 ≈ base 77.97（提交配置）', 'size': 11, 'color': INK, 'bold': True}])
    # decision strip
    rect(s, 0.7, 6.2, 11.92, 0.72, INK)
    txt(s, 1.0, 6.34, 11.4, 0.5, [
        {'t': 'DECISION  ', 'size': 11, 'color': BRIGHT, 'mono': True, 'bold': True,
         'tracking': 1.0},
        {'t': '放弃激进清洗与复杂鲁棒 loss → 把算力押到 完整数据 + 温和 Mixup/CutMix',
         'size': 13, 'color': WHITE}], anchor=MSO_ANCHOR.MIDDLE)


# ===================== SLIDE 5 · 鲁棒损失证伪 (bars) =====================
def s5():
    s = slide()
    header(s, 5, 'LOSS COMPARISON · 鲁棒损失证伪', '4 组实验，0 个超越 CE+LS',
           '40 EP · VAL TOP-1 · RESNET-50 · WEAK AUG')
    bars = [
        ('CE + LS 0.1', 'baseline · 起跑锚点', 66.63, ACCENT, True),
        ('SCE α=.5 β=.5', '与 CE 完全持平 · 零增益', 66.63, GREY3, False),
        ('SCE α=.1 β=1.0', 'β 过大反而退化 −5.97', 60.66, GREY3, False),
        ('GCE q=0.7', 'crashed · 梯度被截断', 26.01, BRIGHT, False),
    ]
    base_y = 6.4          # bottom baseline
    max_h = 3.4           # px height for 100%
    bw = 2.3
    x0 = 1.1
    gap = 0.65
    hline(s, 0.7, base_y, 11.9, GREY2, 1.0)
    for i, (name, note, acc, col, hl) in enumerate(bars):
        bx = x0 + i * (bw + gap)
        bh = max_h * acc / 100.0
        by = base_y - bh
        rect(s, bx, by, bw, bh, col)
        # value on top
        txt(s, bx, by - 0.5, bw, 0.4, {'t': f'{acc:.2f}', 'size': 20,
            'color': (ACCENT if hl else INK), 'bold': True},
            align=PP_ALIGN.CENTER)
        # label under baseline
        txt(s, bx - 0.1, base_y + 0.12, bw + 0.2, 0.35, {'t': name, 'size': 12,
            'color': INK, 'bold': True}, align=PP_ALIGN.CENTER)
        txt(s, bx - 0.1, base_y + 0.5, bw + 0.2, 0.5, {'t': note, 'size': 9.5,
            'color': GREY3}, align=PP_ALIGN.CENTER, line=1.1)
    txt(s, 0.7, 2.0, 11.9, 0.3, {'t': '↑ 柱高 = VAL TOP-1 — 噪声率 5–15% 属弱噪场景，'
        'SCE/GCE 为 40%+ 高噪率设计，假设不成立', 'size': 11.5, 'color': SEC})


# ===================== SLIDE 6 · 消融总览 =====================
def s6():
    s = slide()
    header(s, 6, 'ABLATION · 5 行账单', '核心消融，一图全览',
           'VAL TOP-1 · ALL EXPERIMENTS')
    rows = [
        ('BASELINE_V1', 'full data · RRC + HFlip · CE+LS', '77.58', False),
        ('CLEANED_V1', 'C1+C2 删样本降低覆盖度 · 不作为最终路线', '76.84', False),
        ('FULL_STRONG_AUG', 'Mixup+CutMix+RandAug+Erasing · 100ep 收敛不足', '76.91', False),
        ('FULL_MIXCUT', '最终训练 recipe · best_source=raw', '77.57', False),
        ('FULL_MIXCUT_HFLIP', '同一模型两视角推理 · no ensemble', '77.99', True),
    ]
    y0, rh = 2.5, 0.82
    for i, (name, desc, acc, hl) in enumerate(rows):
        ry = y0 + i * rh
        if hl:
            rect(s, 0.7, ry, 11.92, rh - 0.1, RGBColor(0xF7, 0xEC, 0xEC))
        rect(s, 0.7, ry, 0.06, rh - 0.1, ACCENT if hl else GREY2)
        txt(s, 1.05, ry + 0.06, 4.2, 0.4, {'t': name, 'size': 14,
            'color': INK, 'bold': True, 'mono': True})
        txt(s, 1.05, ry + 0.42, 7.5, 0.3, {'t': desc, 'size': 11, 'color': SEC})
        txt(s, 9.5, ry + 0.05, 3.0, 0.6, {'t': acc, 'size': 26,
            'color': (ACCENT if hl else INK), 'bold': True},
            align=PP_ALIGN.RIGHT)
        if not hl:
            hline(s, 0.7, ry + rh - 0.08, 11.92, GREY2, 0.5)


# ===================== SLIDE 7 · 3 个 Trainer Bug =====================
def s7():
    s = slide()
    header(s, 7, 'ENGINEERING RIGOR · TRAINER BUGS', '3 个 Bug = 3 道方法论',
           '发现 bug 不是失败 — 是把"看不见的实验变量"变成"看得见的方法论"')
    bugs = [
        ('#1', 'LS 静默失效', 'Label Smoothing',
         'mixup 分支跳过 LS，归因不干净', 'mixup 路径显式注入 mixed×(1−ε)+ε/N'),
        ('#2', 'cudnn 配置冲突', 'Performance',
         'determ=True 时 benchmark 不生效，慢 5–15%', 'determ=False · benchmark=True'),
        ('#3', 'resume scheduler', 'Just Fixed · 5/26',
         'LR 反向爬升炸权重，72.51 → 12.64', '--resume-mode finetune 重置调度器'),
    ]
    cw, gap, x0, y0, ch = 3.9, 0.15, 0.7, 2.5, 4.1
    for i, (no, sym, tag, impact, fix) in enumerate(bugs):
        cx = x0 + i * (cw + gap)
        rect(s, cx, y0, cw, ch, WHITE, line=GREY2, lw=0.75)
        txt(s, cx + 0.3, y0 + 0.25, 1.5, 0.5, {'t': no, 'size': 24,
            'color': ACCENT, 'bold': True, 'mono': True})
        txt(s, cx + 0.3, y0 + 0.85, cw - 0.6, 0.4, {'t': sym, 'size': 16,
            'color': INK, 'bold': True})
        txt(s, cx + 0.3, y0 + 1.28, cw - 0.6, 0.3, {'t': tag, 'size': 9.5,
            'color': GREY3, 'mono': True, 'tracking': 0.6})
        txt(s, cx + 0.3, y0 + 1.75, cw - 0.6, 0.3, {'t': 'IMPACT', 'size': 9,
            'color': ACCENT, 'mono': True, 'tracking': 1.0})
        txt(s, cx + 0.3, y0 + 2.02, cw - 0.6, 1.0, {'t': impact, 'size': 11.5,
            'color': SEC}, line=1.2)
        hline(s, cx + 0.3, y0 + 3.05, cw - 0.6, GREY2, 0.75)
        txt(s, cx + 0.3, y0 + 3.18, cw - 0.6, 0.3, {'t': 'FIX', 'size': 9,
            'color': ACCENT, 'mono': True, 'tracking': 1.0})
        txt(s, cx + 0.3, y0 + 3.45, cw - 0.6, 0.9, {'t': fix, 'size': 11,
            'color': INK, 'mono': True}, line=1.2)


# ===================== SLIDE 8 · 推理端增益 =====================
def s8():
    s = slide()
    header(s, 8, 'INFERENCE · 推理端增益', 'HFlip TTA · +0.42 pp · 零额外模型',
           'SINGLE CHECKPOINT · ZERO EXTRA MODEL COST')
    cards = [
        ('01 / HFLIP', 'HFlip TTA', '水平翻转后再前向一次，softmax 平均，对方向偏置鲁棒',
         'softmax(x) ⊕ softmax(flip x)', 'ΔTOP-1', '+0.42 pp'),
        ('02 / SOURCE', 'raw/EMA 双评估', '同一 checkpoint 记录 best_source，自动加载最优权重',
         'raw 77.57  ·  EMA 77.438', 'SOURCE', 'raw'),
        ('03 / MEMORY', 'CPU logits 累加', '10万×1000×4B ≈ 400MB，累加在 CPU 避免 GPU OOM',
         '100k × 1000 · float32', 'DEVICE', 'CPU'),
    ]
    cw, gap, x0, y0, ch = 3.9, 0.15, 0.7, 2.5, 4.0
    for i, (no, name, desc, mech, klab, kval) in enumerate(cards):
        cx = x0 + i * (cw + gap)
        rect(s, cx, y0, cw, ch, WHITE, line=GREY2, lw=0.75)
        rect(s, cx, y0, cw, 0.08, ACCENT)
        txt(s, cx + 0.3, y0 + 0.3, cw - 0.6, 0.3, {'t': no, 'size': 9.5,
            'color': ACCENT, 'mono': True, 'tracking': 0.8})
        txt(s, cx + 0.3, y0 + 0.62, cw - 0.6, 0.5, {'t': name, 'size': 18,
            'color': INK, 'bold': True})
        txt(s, cx + 0.3, y0 + 1.25, cw - 0.6, 1.1, {'t': desc, 'size': 11.5,
            'color': SEC}, line=1.25)
        rect(s, cx + 0.3, y0 + 2.35, cw - 0.6, 0.55, RGBColor(0xF0, 0xF0, 0xEE))
        txt(s, cx + 0.45, y0 + 2.46, cw - 0.85, 0.4, {'t': mech, 'size': 10,
            'color': INK, 'mono': True}, anchor=MSO_ANCHOR.MIDDLE)
        hline(s, cx + 0.3, y0 + 3.15, cw - 0.6, GREY2, 0.75)
        txt(s, cx + 0.3, y0 + 3.28, cw - 0.6, 0.25, {'t': klab, 'size': 9,
            'color': GREY3, 'mono': True, 'tracking': 0.8})
        txt(s, cx + 0.3, y0 + 3.5, cw - 0.6, 0.5, {'t': kval, 'size': 20,
            'color': ACCENT, 'bold': True})


# ===================== SLIDE 9 · Timeline =====================
def s9():
    s = slide()
    header(s, 9, 'TIMELINE · 项目演进', '7 个阶段，每段都有交付',
           '2026 · APR – JUN · USST TEAM')
    rows = [
        ('4/28 – 5/4', '数据准备 · 下载校验 · AutoDL 部署', '128 万张'),
        ('5/5 – 5/11', 'Pipeline 搭建 · 5 种 Loss · StarNet / ResNet-50', '5 losses'),
        ('5/12 – 5/15', '数据问题修复 · val 缺失 · EMA buffer sync', '9 issues'),
        ('5/15 – 5/18', 'Baseline 100ep 收敛 · 主指标已交付', '77.58%'),
        ('5/19 – 5/25', '消融对比 · 4 组 Loss + 2 组增广', '6 runs'),
        ('5/26', 'Trainer 全代码审校 · 3 Bug · raw/EMA dual-eval', '3 bugs'),
        ('5/30 – 6/2', 'full_mixcut 100ep · HFlip TTA 77.99 · Docker', 'FINAL'),
    ]
    y0, rh = 2.4, 0.62
    # vertical spine
    ln = s.shapes.add_connector(2, Inches(2.55), Inches(y0 + 0.1),
                                Inches(2.55), Inches(y0 + len(rows) * rh - 0.2))
    ln.line.color.rgb = GREY2; ln.line.width = Pt(1.5)
    for i, (date, work, tag) in enumerate(rows):
        ry = y0 + i * rh
        last = (i == len(rows) - 1)
        txt(s, 0.5, ry, 1.75, 0.4, {'t': date, 'size': 11.5,
            'color': INK, 'bold': True, 'mono': True}, align=PP_ALIGN.RIGHT)
        dot = rect(s, 2.45, ry + 0.04, 0.2, 0.2, ACCENT if last else WHITE,
                   line=ACCENT, lw=1.5)
        dot.adjustments  # keep ref
        txt(s, 2.95, ry, 7.2, 0.4, {'t': work, 'size': 12, 'color': SEC})
        txt(s, 10.5, ry, 2.1, 0.4, {'t': tag, 'size': 12,
            'color': (ACCENT if last else GREY3), 'bold': last, 'mono': True},
            align=PP_ALIGN.RIGHT)


# ===================== SLIDE 10 · Closing =====================
def s10():
    s = slide()
    # dark closing slide
    rect(s, 0, 0, SW, SH, INK)
    rect(s, 0, 0, 0.22, SH, ACCENT)
    kicker(s, 0.7, 0.7, 'CLOSING · 10 / 10', color=BRIGHT)
    txt(s, 0.66, 1.5, 12, 1.6, [
        {'t': 'Engineering rigor ', 'size': 40, 'color': WHITE, 'bold': True},
        {'t': 'over fancy losses.', 'size': 40, 'color': BRIGHT, 'bold': True}])
    txt(s, 0.68, 2.95, 11.5, 0.5, {'t': '工程严谨性，胜过损失函数堆砌 — '
        '4 组负样本 + 3 个 Bug 复盘 = 1 条干净归因', 'size': 14,
        'color': GREY2})
    hline(s, 0.7, 3.9, 11.9, RGBColor(0x3a, 0x3a, 0x3a), 1.0)
    takeaways = [
        ('已交付', 'Baseline 77.58 · 6 组消融 · 3 Bug 修复 · 单模型 HFlip TTA · Docker'),
        ('最终路线', 'full_mixcut plain 77.57 · HFlip TTA 77.99 · best_source=raw'),
        ('交付物', 'result.csv · Dockerfile · 技术方案 · 答辩 deck'),
    ]
    for i, (h, d) in enumerate(takeaways):
        ty = 4.2 + i * 0.78
        txt(s, 0.7, ty, 1.0, 0.4, {'t': f'{i+1:02d}', 'size': 16,
            'color': BRIGHT, 'bold': True, 'mono': True})
        txt(s, 1.7, ty, 2.0, 0.4, {'t': h, 'size': 14, 'color': WHITE,
            'bold': True})
        txt(s, 3.8, ty + 0.02, 8.6, 0.4, {'t': d, 'size': 11.5,
            'color': GREY2})
    hline(s, 0.7, 6.7, 11.9, RGBColor(0x3a, 0x3a, 0x3a), 1.0)
    txt(s, 0.7, 6.85, 8, 0.4, {'t': 'USST · D Group · Chen   |   2026.06.14',
        'size': 11, 'color': GREY3, 'mono': True})
    txt(s, 7.5, 6.85, 5.1, 0.4, {'t': '谢谢 · 欢迎提问 · Q & A', 'size': 13,
        'color': BRIGHT, 'bold': True}, align=PP_ALIGN.RIGHT)


def main():
    for fn in (s1, s2, s3, s4, s5, s6, s7, s8, s9, s10):
        fn()
    out = r'a:\VScode\Code\Projects\embodied_ai_contest\submit\tech_report\答辩PPT.pptx'
    prs.save(out)
    print('saved', out, '·', len(prs.slides._sldIdLst), 'slides')


if __name__ == '__main__':
    main()









